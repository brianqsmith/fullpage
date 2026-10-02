from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1];source=ROOT/'docs/01_FUNCTIONAL_REQUIREMENTS.md'
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.widow_control=True
normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.10;normal.paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(27);doc.styles['Title'].paragraph_format.space_after=Pt(14)
for name,size in [('Heading 1',16),('Heading 2',12)]:
 st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(15);st.paragraph_format.space_after=Pt(7);st.paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('fullpage 1.4.0  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def runs(p,text):
 for chunk in re.split(r'(\*\*.*?\*\*|`[^`]+`)',text):
  if not chunk:continue
  r=p.add_run(chunk[2:-2] if chunk.startswith('**') else chunk[1:-1] if chunk.startswith('`') else chunk)
  if chunk.startswith('**'):r.bold=True
  if chunk.startswith('`'):r.font.name='Consolas';r.font.size=Pt(9)
lines=source.read_text().splitlines();i=0
while i<len(lines):
 line=lines[i].strip();i+=1
 if not line:continue
 if line.startswith('|'):
  rows=[line]
  while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
  rows=[x for x in rows if not re.match(r'^\|[\s:|\-]+$',x)]
  table=doc.add_table(rows=0,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
  widths=[2.0,1.15,3.62]
  for col,width in zip(table.columns,widths):col.width=Inches(width)
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
  table._tbl.tblPr.append(borders)
  for idx,rowtext in enumerate(rows):
   cells=table.add_row().cells
   for col,text in enumerate(rowtext.strip('|').split('|')):
    cell=cells[col];cell.width=Inches(widths[col]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5);p.paragraph_format.line_spacing=1.05;runs(p,text.strip())
    tcPr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
    for side in ['top','left','bottom','right']:
     item=OxmlElement('w:'+side);item.set(qn('w:w'),'90');item.set(qn('w:type'),'dxa');mar.append(item)
    tcPr.append(mar)
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E4EBF0' if idx==0 else ('F5F7F8' if idx%2==0 else 'FFFFFF'));tcPr.append(shade)
    for run in p.runs:run.font.size=Pt(9);run.bold=idx==0
   trPr=table.rows[-1]._tr.get_or_add_trPr();avoid=OxmlElement('w:cantSplit');trPr.append(avoid)
   if idx==0:trPr.append(OxmlElement('w:tblHeader'))
  doc.add_paragraph();continue
 if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
 elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
 elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
 elif line.startswith('- '):runs(doc.add_paragraph(style='List Bullet'),line[2:])
 else:runs(doc.add_paragraph(),line)
doc.core_properties.title='fullpage functional requirements';doc.core_properties.subject='Developer takeover specification for version 1.4.0';doc.core_properties.author='fullpage project';doc.core_properties.keywords='Firefox Manifest V3 functional requirements developer handoff'
for element in [doc._element,doc.styles.element]:
 for border in element.xpath('.//w:pBdr'):border.getparent().remove(border)
out=ROOT/'docs/01_FUNCTIONAL_REQUIREMENTS.docx';doc.save(out);print(out)
